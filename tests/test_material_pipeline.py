import sys
import unittest
import re
from io import BytesIO
from pathlib import Path

from docx import Document


PROJECT_DIR = Path(__file__).resolve().parents[1]
WORKSPACE_DIR = PROJECT_DIR.parent
if str(PROJECT_DIR) not in sys.path:
    sys.path.insert(0, str(PROJECT_DIR))

from services.knowledge_json_extractor import build_knowledge_json_prompt, extract_knowledge_points_as_drafts
from services.material_cleaner import clean_material_for_extraction
from services.material_router import route_material_input
from services.pdf_outline_service import extract_outline_candidates
from services.pdf_text_service import extract_pdf_text
from schemas.material_schema import MaterialResult


class MaterialPipelineTests(unittest.TestCase):
    def test_material_result_can_restore_older_persisted_payload(self):
        result = MaterialResult.from_dict(
            {
                "source_type": "pdf",
                "process_method": "pdf_text_extract",
                "extracted_text": "已确认的资料文本",
                "confidence": 0.92,
                "future_field": "ignored",
            }
        )
        self.assertEqual(result.source_type, "pdf")
        self.assertEqual(result.extracted_text, "已确认的资料文本")
        self.assertEqual(result.warnings, [])

        damaged = MaterialResult.from_dict(
            {
                "source_type": "unknown",
                "process_method": None,
                "extracted_text": 123,
                "confidence": "not-a-number",
                "warnings": "旧版警告",
                "clean_report": [],
            }
        )
        self.assertEqual(damaged.source_type, "pasted_text")
        self.assertEqual(damaged.confidence, 0.0)
        self.assertEqual(damaged.warnings, ["旧版警告"])
        self.assertEqual(damaged.clean_report, {})

        outline = MaterialResult.from_dict(
            {
                "source_type": "pdf",
                "process_method": "pdf_outline_ai",
                "extracted_text": "第一章 绪论",
                "confidence": 0.6,
            }
        )
        self.assertEqual(outline.process_method, "pdf_outline_ai")

    def test_subject_guidance_is_added_without_weakening_source_grounding(self):
        prompt = build_knowledge_json_prompt(
            "题1：肺炎链球菌是社区获得性肺炎的常见病原体。",
            subject="医学考研",
            chapter_name="内科学",
            extraction_guidance="优先识别诊断依据、鉴别诊断和用药禁忌。",
        )
        self.assertIn("诊断依据、鉴别诊断和用药禁忌", prompt)
        self.assertIn("只基于用户提供的资料内容", prompt)
        self.assertIn("不能突破原文依据", prompt)
        self.assertIn("学习心得、备考经验", prompt)
        self.assertIn("严禁保存广告、扫码咨询", prompt)

    def test_outline_prompt_marks_only_generated_content_as_ai_expansion(self):
        prompt = build_knowledge_json_prompt(
            "=== 第3页 ===\n- 第一章 计算机系统概述",
            subject="408综合",
            outline_mode=True,
        )
        self.assertIn("输入内容是从图片型 PDF 抽样识别出的目录/章节提纲", prompt)
        self.assertIn('原文是题目，应设置 knowledge_type="题目"', prompt)
        self.assertIn("不得把封面、书名、作者、出版社", prompt)

        drafts, _warnings = extract_knowledge_points_as_drafts(
            "=== 第3页 ===\n- 第一章 计算机系统概述",
            subject="408综合",
            outline_mode=True,
            llm_callable=lambda _prompt: """[
                {
                    "knowledge_name": "计算机系统层次结构",
                    "core_definition": "计算机系统由硬件和软件共同构成。",
                    "source_text": "第一章 计算机系统概述",
                    "is_ai_expansion": false
                }
            ]""",
        )
        self.assertTrue(drafts)
        self.assertTrue(drafts[0].is_ai_expansion)
        self.assertIn("基于提纲AI发散，需核对教材", drafts[0].uncertainty_note)

    def test_large_syllabus_outline_is_read_in_one_whole_document_pass(self):
        outline = "【考试大纲提纲（原文层级抽取）】\n\n" + "\n\n".join(
            f"=== 第{page}页 ===\n"
            + "\n".join(
                f"[二级] （{index}）知识点{page}-{index}的基本概念、实现原理与典型考法"
                for index in range(1, 21)
            )
            for page in range(1, 13)
        )
        prompts = []

        def fake_llm(prompt):
            prompts.append(prompt)
            return '["K001", "K240"]'

        drafts, _warnings = extract_knowledge_points_as_drafts(
            outline,
            subject="408综合",
            max_points=48,
            outline_mode=True,
            llm_callable=fake_llm,
        )

        self.assertEqual(len(prompts), 1)
        self.assertIn("知识点1-1", prompts[0])
        self.assertIn("知识点12-20", prompts[0])
        self.assertIn("[K001]", prompts[0])
        self.assertIn("[K240]", prompts[0])
        self.assertEqual(len(drafts), 48)
        self.assertEqual(drafts[0].knowledge_name, "知识点1-1的基本概念、实现原理与典型考法")
        self.assertEqual(drafts[1].knowledge_name, "知识点12-20的基本概念、实现原理与典型考法")
        self.assertTrue(all(not draft.is_ai_expansion for draft in drafts))
        self.assertTrue(all(not draft.core_definition for draft in drafts))
        self.assertTrue(all(draft.knowledge_type == "大纲知识点" for draft in drafts))

    def test_syllabus_copy_is_not_mislabeled_as_ai_expansion_or_question(self):
        drafts, _warnings = extract_knowledge_points_as_drafts(
            "【考试大纲提纲（原文层级抽取）】\n=== 第2页 ===\n[三级] 1.顺序存储",
            subject="408综合",
            max_points=6,
            outline_mode=True,
            llm_callable=lambda _prompt: '["K001"]',
        )

        self.assertEqual(len(drafts), 1)
        self.assertFalse(drafts[0].is_ai_expansion)
        self.assertEqual(drafts[0].knowledge_type, "大纲知识点")

    def test_syllabus_parent_headings_are_not_saved_as_knowledge_points(self):
        outline = """【考试大纲提纲（原文层级抽取）】
=== 第1页 ===
[科目] 计算机组成原理
[一级] 一、计算机系统概述
[二级] （一）计算机系统层次结构
[三级] 1.计算机硬件的基本组成
[三级] 2.计算机软件的分类
[二级] （二）计算机系统的性能指标
"""
        prompts = []

        def fake_llm(prompt):
            prompts.append(prompt)
            return '["K001"]'

        drafts, _warnings = extract_knowledge_points_as_drafts(
            outline,
            subject="408",
            max_points=10,
            outline_mode=True,
            llm_callable=fake_llm,
        )

        names = [draft.knowledge_name for draft in drafts]
        self.assertEqual(len(prompts), 1)
        self.assertNotIn("计算机系统层次结构", names)
        self.assertIn("计算机硬件的基本组成", names)
        self.assertIn("计算机软件的分类", names)
        self.assertIn("计算机系统的性能指标", names)

    def test_outline_fallback_filters_cover_metadata_and_keeps_questions(self):
        outline = """
=== 第1页 ===
- 金榜时代
- GLISTIMELTE
- 编著◎武忠祥
- 中国农业出版社
=== 第4页 ===
- 第一章
- 函数极限连续
- 第二章 导数与微分
=== 第6页 ===
- 第一章 函数极限连续
题1：设函数 f(x) 满足条件，判断其奇偶性。 (A) 偶函数 (B) 奇函数
"""
        drafts, _warnings = extract_knowledge_points_as_drafts(
            outline,
            subject="数二",
            chapter_name="高数严选题",
            outline_mode=True,
            llm_callable=lambda _prompt: (_ for _ in ()).throw(RuntimeError("mock llm failure")),
        )
        payloads = [draft.__dict__ for draft in drafts]
        names = [item["knowledge_name"] for item in payloads]
        self.assertNotIn("金榜时代", names)
        self.assertNotIn("GLISTIMELTE", names)
        self.assertNotIn("中国农业出版社", names)
        self.assertTrue(any(item["knowledge_type"] == "题目" for item in payloads))
        question = next(item for item in payloads if item["knowledge_type"] == "题目")
        self.assertIn("函数性质判断", question["knowledge_name"])
        self.assertEqual(question["exam_question_styles"], ["选择题"])
        self.assertFalse(question["is_ai_expansion"])

    def test_outline_candidate_filter_rejects_publication_metadata(self):
        candidates = extract_outline_candidates(
            """
金榜时代
编著 武忠祥
中国农业出版社
第一章 函数极限连续
题1：判断函数的奇偶性 (A) 偶函数 (B) 奇函数
"""
        )
        self.assertEqual(candidates[0], "第一章 函数极限连续")
        self.assertTrue(any(item.startswith("题1") for item in candidates))
        self.assertFalse(any("出版社" in item or "编著" in item for item in candidates))

    def test_cleaner_removes_promo_lines_but_preserves_question_context(self):
        raw_text = """
=== 第1页 ===
1. 下列关于微信消息队列机制的说法，正确的是
A. 生产者发送后立即返回
微信公众号 计算机与软件考研
关注微信公众号 计算机与软件考研
"""
        result = clean_material_for_extraction(raw_text)
        self.assertIn("微信消息队列机制", result.cleaned_text)
        self.assertNotIn("关注微信公众号", result.cleaned_text)
        self.assertGreaterEqual(result.removed_noise_lines, 1)

    def test_cleaner_keeps_experience_but_removes_embedded_consultation_ad(self):
        raw_text = (
            "=== 第1页 ===\n"
            "更多计算机考研资料和信息，请扫码咨询>>> "
            "考场时间应控制在2小时45分钟，最后15分钟用于检查。\n"
            "考研经验：遇到明显异常的题目应先跳过，避免打乱整体节奏。"
        )
        result = clean_material_for_extraction(raw_text)
        self.assertNotIn("扫码咨询", result.cleaned_text)
        self.assertIn("2小时45分钟", result.cleaned_text)
        self.assertIn("考研经验", result.cleaned_text)

    def test_route_material_input_returns_clean_report_for_text(self):
        result = route_material_input(
            pasted_text="=== 第1页 ===\n1. 栈的特点是先进后出还是后进先出？\nA. 先进后出",
        )
        self.assertEqual(result.process_method, "pasted_text")
        self.assertTrue(result.extracted_text)
        self.assertIn("cleaned_text_length", result.clean_report)
        self.assertGreaterEqual(result.clean_report.get("page_markers", 0), 1)

    def test_all_advertised_non_pdf_formats_have_working_routes(self):
        for suffix in ("txt", "md"):
            result = route_material_input(
                file_name=f"sample.{suffix}",
                file_bytes="函数极限与连续是高等数学的基础内容。".encode("utf-8"),
            )
            self.assertEqual(result.process_method, "pasted_text")
            self.assertIn("函数极限", result.extracted_text)

        for suffix in ("png", "jpg", "jpeg"):
            result = route_material_input(
                file_name=f"sample.{suffix}",
                file_bytes=b"mock-image",
                image_ocr_fn=lambda _bytes: "题1：函数极限的计算方法。",
            )
            self.assertEqual(result.process_method, "image_ocr")
            self.assertIn("函数极限", result.extracted_text)

        buffer = BytesIO()
        document = Document()
        document.add_heading("中国近现代史", level=1)
        document.add_paragraph("鸦片战争是中国近代史的重要开端。")
        document.add_paragraph("洋务运动兴办近代工业。", style="List Bullet")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "时间"
        table.cell(0, 1).text = "事件"
        table.cell(1, 0).text = "1840年"
        table.cell(1, 1).text = "鸦片战争"
        document.save(buffer)

        result = route_material_input(
            file_name="中国近现代史.docx",
            file_bytes=buffer.getvalue(),
        )
        self.assertEqual(result.source_type, "docx")
        self.assertEqual(result.process_method, "docx_text_extract")
        self.assertIn("【1级标题】中国近现代史", result.extracted_text)
        self.assertIn("- 洋务运动兴办近代工业。", result.extracted_text)
        self.assertIn("1840年 | 鸦片战争", result.extracted_text)
        self.assertEqual(result.ocr_report.get("table_row_count"), 2)

    def test_invalid_docx_is_rejected_with_actionable_message(self):
        with self.assertRaisesRegex(ValueError, "另存为"):
            route_material_input(
                file_name="损坏资料.docx",
                file_bytes=b"not-a-docx",
            )

    def test_docx_material_result_round_trips_without_type_loss(self):
        restored = MaterialResult.from_dict(
            {
                "source_type": "docx",
                "process_method": "docx_text_extract",
                "extracted_text": "第一章 中国古代史",
                "confidence": 0.97,
            }
        )
        self.assertEqual(restored.source_type, "docx")
        self.assertEqual(restored.process_method, "docx_text_extract")

    def test_cleaner_splits_fullwidth_question_numbers_and_keeps_other_pages(self):
        raw_text = """
=== 第1页 ===
22. 下列关于DMA方式的叙述中，正确的是。
A. DMA 传送前由设备驱动程序设置传送参数

=== 第4页 ===
题26：下列选项中，可用于文件系统管理空闲磁盘块的数据结构是。 I．位图 II．索引结点 III．空闲磁盘块链 IV．文件分配表(FAT) A. 仅I、II B. 仅I、III、IV C. 仅I、III D. 仅II、III、IV 27．系统采用二级反馈队列调度算法进行进程调度。就绪队列Q1采用时间片轮转调度算法。
"""
        result = clean_material_for_extraction(raw_text)
        self.assertIn("=== 第1页 ===", result.cleaned_text)
        self.assertIn("题22：", result.cleaned_text)
        self.assertIn("=== 第4页 ===", result.cleaned_text)
        self.assertIn("题26：", result.cleaned_text)
        self.assertIn("题27：", result.cleaned_text)
        self.assertGreaterEqual(result.question_blocks, 3)

    def test_cleaner_preserves_pipeline_table_layout(self):
        raw_text = """
=== 第11页 ===
41. 某五段式流水线如下图所示：
时间单元
指令
I1
IF
ID
EX
M
WB
I2
IF
ID
EX
M
WB
读寄存器内容，所以 I3 的 ID 段被阻塞。
"""
        result = clean_material_for_extraction(raw_text)
        self.assertIn("时间单元\n指令\nI1\nIF\nID\nEX\nM\nWB", result.cleaned_text)
        self.assertIn("I3 的 ID 段被阻塞", result.cleaned_text)

    def test_pdf_cleaning_improves_real_408_sample(self):
        sample_pdf = WORKSPACE_DIR / "kaoyan-assistant" / "data" / "test_materials" / "408" / "2025-408真题.pdf"
        if not sample_pdf.exists():
            self.skipTest(f"sample pdf not found: {sample_pdf}")

        pdf_data = extract_pdf_text(sample_pdf)
        cleaned = clean_material_for_extraction(pdf_data["text"])
        self.assertIn("题1", cleaned.cleaned_text)
        self.assertTrue(
            cleaned.removed_noise_lines > 0
            or cleaned.removed_inline_noise > 0
            or len(cleaned.cleaned_text) != len(pdf_data["text"])
        )
        self.assertGreaterEqual(cleaned.question_blocks, 10)

    def test_pdf_text_extract_keeps_long_real_sample_untruncated(self):
        sample_pdf = WORKSPACE_DIR / "kaoyan-assistant" / "data" / "test_materials" / "408" / "2025-408真题.pdf"
        if not sample_pdf.exists():
            self.skipTest(f"sample pdf not found: {sample_pdf}")

        pdf_data = extract_pdf_text(sample_pdf)
        self.assertIn("=== 第2页 ===", pdf_data["text"])
        self.assertGreater(len(pdf_data["text"]), 8000)

    def test_long_draft_extraction_uses_later_pages_instead_of_truncating(self):
        text = "\n\n".join(
            [
                "=== 第1页 ===\n题1：缓存命中率影响平均访存时间。" + (" 补充说明" * 420),
                "=== 第6页 ===\n题2：分页系统需要页表与快表配合。" + (" 补充说明" * 420),
                "=== 第11页 ===\n题3：五段式流水线中，读寄存器内容，所以 I3 的 ID 段被阻塞。",
            ]
        )
        drafts, warnings = extract_knowledge_points_as_drafts(
            text,
            subject="408综合",
            chapter_name="组成原理",
            llm_callable=lambda _prompt: (_ for _ in ()).throw(RuntimeError("mock llm failure")),
        )
        self.assertTrue(any("分段抽取" in warning for warning in warnings))
        self.assertTrue(any(getattr(draft, "source_page", "") == "第11页" for draft in drafts))

    def test_very_long_draft_extraction_evenly_covers_last_page_within_budget(self):
        text = "\n\n".join(
            f"=== 第{page}页 ===\n题{page}：第{page}页的核心机制用于验证长文档覆盖。"
            + ("该页包含独立且可复习的详细说明。" * 120)
            for page in range(1, 31)
        )

        drafts, warnings = extract_knowledge_points_as_drafts(
            text,
            subject="408综合",
            chapter_name="长文档覆盖测试",
            max_points=6,
            llm_callable=lambda _prompt: (_ for _ in ()).throw(RuntimeError("mock llm failure")),
        )

        self.assertLessEqual(len(drafts), 6)
        self.assertTrue(any(getattr(draft, "source_page", "") == "第30页" for draft in drafts))
        self.assertTrue(any("均匀选取" in warning for warning in warnings))
        self.assertTrue(any("覆盖首段、中段和末段" in warning for warning in warnings))
        self.assertTrue(any("未逐段调用模型" in warning for warning in warnings))

    def test_image_dominant_pdf_forces_ocr_fallback(self):
        sample_pdf = WORKSPACE_DIR / "kaoyan-assistant" / "data" / "test_materials" / "820" / "电子科技大学-820-2012-真题.pdf"
        if not sample_pdf.exists():
            self.skipTest(f"sample pdf not found: {sample_pdf}")

        result = route_material_input(
            file_name=sample_pdf.name,
            file_path=str(sample_pdf),
            pdf_ocr_available=True,
            pdf_ocr_fn=lambda _path: (
                "=== 第1页 ===\n题1：操作系统的主要功能包括处理机管理、存储器管理和文件管理。",
                {"primary_engine": "RapidOCR", "pages_processed": 1},
            ),
        )
        self.assertEqual(result.process_method, "pdf_ocr")
        self.assertTrue(result.pdf_diagnostics.get("needs_ocr"))
        self.assertEqual(result.ocr_report.get("primary_engine"), "RapidOCR")
        self.assertIn("题1", result.extracted_text)
        self.assertTrue(any("重复水印" in warning or "图片主导" in warning for warning in result.warnings))

    def test_image_dominant_pdf_prefers_outline_mode_when_available(self):
        sample_pdf = WORKSPACE_DIR / "kaoyan-assistant" / "data" / "test_materials" / "820" / "电子科技大学-820-2012-真题.pdf"
        if not sample_pdf.exists():
            self.skipTest(f"sample pdf not found: {sample_pdf}")

        result = route_material_input(
            file_name=sample_pdf.name,
            file_path=str(sample_pdf),
            pdf_ocr_available=True,
            pdf_outline_fn=lambda _path: (
                "=== 第1页 ===\n- 第一章 操作系统概述",
                {"mode": "outline", "sampled_pages": [1]},
            ),
            pdf_ocr_fn=lambda _path: (_ for _ in ()).throw(AssertionError("不应执行全量 OCR")),
        )
        self.assertEqual(result.process_method, "pdf_outline_ai")
        self.assertEqual(result.ocr_report.get("mode"), "outline")
        self.assertIn("第一章", result.extracted_text)
        self.assertTrue(any("必须结合教材人工核对" in warning for warning in result.warnings))

    def test_regular_text_pdf_keeps_direct_extraction(self):
        sample_pdf = WORKSPACE_DIR / "kaoyan-assistant" / "data" / "test_materials" / "408" / "2025-408真题.pdf"
        if not sample_pdf.exists():
            self.skipTest(f"sample pdf not found: {sample_pdf}")

        result = route_material_input(
            file_name=sample_pdf.name,
            file_path=str(sample_pdf),
            pdf_ocr_available=True,
            pdf_ocr_fn=lambda _path: "unused",
        )
        self.assertEqual(result.process_method, "pdf_text_extract")
        self.assertFalse(result.pdf_diagnostics.get("needs_ocr"))

    def test_draft_extraction_has_local_fallback(self):
        text = "=== 第1页 ===\n题1：栈是一种只允许在一端进行插入和删除的线性表。常见考法是判断出栈序列是否合法。"
        drafts, warnings = extract_knowledge_points_as_drafts(
            text,
            subject="408综合",
            chapter_name="数据结构",
            llm_callable=lambda _prompt: (_ for _ in ()).throw(RuntimeError("mock llm failure")),
        )
        self.assertTrue(drafts)
        self.assertTrue(any("本地兜底抽取" in warning for warning in warnings))
        self.assertEqual(drafts[0].subject, "408综合")

    def test_medical_subject_draft_extraction_keeps_subject_label(self):
        text = "=== 第1页 ===\n题1：社区获得性肺炎常见病原体包括肺炎链球菌。治疗需要结合患者年龄、基础疾病和影像学表现综合判断。"
        drafts, warnings = extract_knowledge_points_as_drafts(
            text,
            subject="医学考研",
            chapter_name="内科学",
            llm_callable=lambda _prompt: (_ for _ in ()).throw(RuntimeError("mock llm failure")),
        )
        self.assertTrue(drafts)
        self.assertTrue(any("本地兜底抽取" in warning for warning in warnings))
        self.assertEqual(drafts[0].subject, "医学考研")

    def test_draft_extraction_reports_progress_stages(self):
        events = []
        text = "\n\n".join(
            [
                "=== 第1页 ===\n题1：缓存命中率影响平均访存时间。" + (" 补充说明" * 420),
                "=== 第6页 ===\n题2：分页系统需要页表与快表配合。" + (" 补充说明" * 420),
                "=== 第11页 ===\n题3：五段式流水线中，读寄存器内容，所以 I3 的 ID 段被阻塞。",
            ]
        )
        drafts, warnings = extract_knowledge_points_as_drafts(
            text,
            subject="408综合",
            chapter_name="组成原理",
            llm_callable=lambda _prompt: (_ for _ in ()).throw(RuntimeError("mock llm failure")),
            progress_callback=lambda current, total, message: events.append((current, total, message)),
        )
        self.assertTrue(drafts)
        self.assertTrue(events)
        self.assertIn("正在整理待抽取文本", events[0][2])
        self.assertIn("抽取完成", events[-1][2])
        self.assertEqual(events[-1][0], events[-1][1])


if __name__ == "__main__":
    unittest.main()
